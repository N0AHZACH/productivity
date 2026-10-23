"""Deterministic DocumentSpec to Word rendering."""

from datetime import date
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.shared import Cm, Pt

from app.core.errors import RenderError
from app.models.document import (
    Block,
    BulletList,
    Callout,
    DocumentSpec,
    Heading,
    NumberedList,
    PageBreak,
    Paragraph,
    Section,
    Table,
)
from app.render.styles import (
    add_page_count_field,
    add_page_number_field,
    add_toc_field,
    configure_document_styles,
    configure_page,
    enable_field_updates,
    format_run,
    hex_to_rgb,
    repeat_table_header,
    set_cell_background,
    set_cell_borders,
    set_cell_margins,
    set_paragraph_bottom_border,
)
from app.render.themes import DEFAULT_THEME, Theme

_CALLOUT_COLORS: dict[str, tuple[str, str]] = {
    "info": ("2E75B6", "EAF2FB"),
    "warning": ("C55A11", "FFF4E5"),
    "risk": ("C00000", "FCEAEA"),
    "decision": ("548235", "EAF4EA"),
}
_CALLOUT_LABELS = {
    "info": "NOTE",
    "warning": "WARNING",
    "risk": "RISK",
    "decision": "DECISION",
}


class DocxRenderer:
    """Render a validated DocumentSpec into a styled .docx file."""

    def __init__(self, theme: Theme = DEFAULT_THEME) -> None:
        self.theme = theme

    def render(self, spec: DocumentSpec) -> bytes:
        """Render a document to bytes without touching the filesystem."""
        document = Document()
        configure_document_styles(document, self.theme)
        page = document.sections[0]
        configure_page(page)
        self._configure_page_furniture(document, spec)
        self._set_core_properties(document, spec)

        self._render_cover(document, spec)
        self._render_toc(document)
        self._render_assumptions(document, spec)
        for section in spec.sections:
            self._render_section(document, section)
        enable_field_updates(document)

        stream = BytesIO()
        document.save(stream)
        return stream.getvalue()

    def render_to(self, spec: DocumentSpec, path: str | Path) -> Path:
        """Render a document and write it to a path."""
        target = Path(path)
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(self.render(spec))
        return target

    def _set_core_properties(self, document: DocumentObject, spec: DocumentSpec) -> None:
        properties = document.core_properties
        properties.title = spec.title or "Untitled Document"
        properties.subject = spec.subtitle or ""
        properties.author = spec.metadata.get("author", self.theme.footer_text or "Productivity")
        properties.keywords = spec.doc_type.value

    def _configure_page_furniture(self, document: DocumentObject, spec: DocumentSpec) -> None:
        section = document.sections[0]
        section.different_first_page_header_footer = True
        header = section.header
        header.is_linked_to_previous = False
        header_paragraph = header.paragraphs[0]
        header_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if spec.title:
            header_run = header_paragraph.add_run(spec.title)
            header_run.font.size = Pt(8.5)
            header_run.font.color.rgb = hex_to_rgb("808080")
            header_run.font.name = self.theme.font_body

        footer = section.footer
        footer.is_linked_to_previous = False
        footer_paragraph = footer.paragraphs[0]
        footer_paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(16.6), WD_TAB_ALIGNMENT.RIGHT)
        if self.theme.footer_text:
            footer_run = footer_paragraph.add_run(self.theme.footer_text)
            footer_run.font.size = Pt(8.5)
            footer_run.font.color.rgb = hex_to_rgb("808080")
        footer_paragraph.add_run("\t")
        page_label = footer_paragraph.add_run("Page ")
        page_label.font.size = Pt(8.5)
        page_label.font.color.rgb = hex_to_rgb("808080")
        add_page_number_field(footer_paragraph)
        of_label = footer_paragraph.add_run(" of ")
        of_label.font.size = Pt(8.5)
        of_label.font.color.rgb = hex_to_rgb("808080")
        add_page_count_field(footer_paragraph)
        for run in footer_paragraph.runs:
            run.font.size = Pt(8.5)
            run.font.color.rgb = hex_to_rgb("808080")

    def _render_cover(self, document: DocumentObject, spec: DocumentSpec) -> None:
        logo = self.theme.logo_path
        if logo:
            logo_path = Path(logo)
            if logo_path.is_file():
                document.add_picture(str(logo_path), width=Cm(3.4))
                document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.LEFT

        spacer = document.add_paragraph()
        spacer.paragraph_format.space_before = Cm(4.5)

        title = document.add_paragraph(style="Title")
        title.add_run(spec.title or "Untitled Document")
        if spec.subtitle:
            subtitle = document.add_paragraph(style="Subtitle")
            subtitle.add_run(spec.subtitle)

        metadata = spec.metadata
        rows = [
            ("Author", metadata.get("author", "Not specified")),
            ("Date", metadata.get("date", date.today().isoformat())),
            ("Version", metadata.get("version", "1.0")),
            ("Document type", spec.doc_type.value.replace("_", " ").title()),
        ]
        table = document.add_table(rows=0, cols=2)
        table.autofit = False
        for label, value in rows:
            cells = table.add_row().cells
            cells[0].width = Cm(3.8)
            cells[1].width = Cm(12.6)
            label_paragraph = cells[0].paragraphs[0]
            format_run(
                label_paragraph,
                label.upper(),
                bold=True,
                color=self.theme.accent_hex,
                size_pt=8,
                font=self.theme.font_body,
            )
            value_paragraph = cells[1].paragraphs[0]
            format_run(
                value_paragraph,
                value or "Not specified",
                color="404040",
                size_pt=9.5,
                font=self.theme.font_body,
            )
            for cell in cells:
                set_cell_borders(cell, bottom=(self.theme.accent_light_hex, 4))

        rule = document.add_paragraph()
        rule.paragraph_format.space_before = Pt(20)
        rule.paragraph_format.space_after = Pt(0)
        set_paragraph_bottom_border(rule, self.theme.accent_hex, size=16)
        document.add_page_break()

    def _render_toc(self, document: DocumentObject) -> None:
        heading = document.add_paragraph()
        format_run(
            heading,
            "Table of Contents",
            bold=True,
            color=self.theme.accent_hex,
            size_pt=self.theme.h1_pt,
            font=self.theme.font_heading,
        )
        heading.paragraph_format.space_after = Pt(16)
        toc_paragraph = document.add_paragraph()
        add_toc_field(toc_paragraph)
        document.add_page_break()

    def _render_assumptions(self, document: DocumentObject, spec: DocumentSpec) -> None:
        if not spec.assumptions:
            return
        document.add_paragraph("Assumptions and Open Questions", style="Heading 1")
        table = document.add_table(rows=1, cols=1)
        table.autofit = False
        cell = table.cell(0, 0)
        cell.width = Cm(16.6)
        set_cell_background(cell, _CALLOUT_COLORS["info"][1])
        set_cell_borders(
            cell,
            left=(_CALLOUT_COLORS["info"][0], 24),
            top=(self.theme.accent_light_hex, 4),
            bottom=(self.theme.accent_light_hex, 4),
            right=(self.theme.accent_light_hex, 4),
        )
        set_cell_margins(cell, top_cm=0.25, bottom_cm=0.2, left_cm=0.35, right_cm=0.25)
        first = cell.paragraphs[0]
        format_run(
            first,
            "The following assumptions were made where the request was incomplete.",
            italic=True,
            color="404040",
            size_pt=9.5,
        )
        for assumption in spec.assumptions:
            if not assumption.strip():
                continue
            bullet = cell.add_paragraph(style="List Bullet")
            bullet.paragraph_format.space_after = Pt(2)
            bullet.add_run(assumption)

    def _render_section(self, document: DocumentObject, section: Section) -> None:
        if section.heading:
            document.add_paragraph(section.heading, style="Heading 1")
        for index, block in enumerate(section.blocks):
            try:
                self._render_block(document, block)
            except RenderError:
                raise
            except Exception as exc:
                raise RenderError(
                    f"Failed to render block {index} in section {section.id!r}: {exc}"
                ) from exc

    def _render_block(self, document: DocumentObject, block: Block) -> None:
        if isinstance(block, Heading):
            if block.text.strip():
                style = f"Heading {min(block.level + 1, 4)}"
                document.add_paragraph(block.text, style=style)
            return
        if isinstance(block, Paragraph):
            if block.text.strip():
                document.add_paragraph(block.text)
            return
        if isinstance(block, BulletList):
            self._render_bullets(document, block)
            return
        if isinstance(block, NumberedList):
            self._render_numbered(document, block)
            return
        if isinstance(block, Table):
            self._render_table(document, block)
            return
        if isinstance(block, Callout):
            self._render_callout(document, block)
            return
        if isinstance(block, PageBreak):
            document.add_page_break()
            return
        raise RenderError(f"Unsupported block type: {type(block).__name__}")

    def _render_bullets(self, document: DocumentObject, block: BulletList) -> None:
        for item in block.items:
            if not item.strip():
                continue
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.space_after = Pt(3)
            paragraph.add_run(item)

    def _render_numbered(self, document: DocumentObject, block: NumberedList) -> None:
        items = [item for item in block.items if item.strip()]
        for number, item in enumerate(items, start=1):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.75)
            paragraph.paragraph_format.first_line_indent = Cm(-0.55)
            paragraph.paragraph_format.space_after = Pt(3)
            paragraph.add_run(f"{number}.\t{item}")

    def _render_table(self, document: DocumentObject, block: Table) -> None:
        if block.caption:
            caption = document.add_paragraph(style="Caption")
            caption.paragraph_format.keep_with_next = True
            caption.add_run(block.caption)
        table = document.add_table(rows=1, cols=len(block.columns))
        table.style = "Table Grid"
        table.autofit = False
        width_cm = 16.6
        widths = self._column_widths(block, width_cm)
        header_cells = table.rows[0].cells
        for index, column in enumerate(block.columns):
            cell = header_cells[index]
            cell.width = Cm(widths[index])
            set_cell_background(cell, self.theme.accent_hex)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            format_run(
                paragraph,
                column,
                bold=True,
                color="FFFFFF",
                size_pt=self.theme.body_pt - 1.5,
                font=self.theme.font_body,
            )
        repeat_table_header(table.rows[0])

        for row_index, row in enumerate(block.rows):
            cells = table.add_row().cells
            for column_index, value in enumerate(row):
                cell = cells[column_index]
                cell.width = Cm(widths[column_index])
                if row_index % 2 == 1:
                    set_cell_background(cell, self.theme.accent_light_hex)
                paragraph = cell.paragraphs[0]
                paragraph.paragraph_format.space_after = Pt(0)
                format_run(
                    paragraph,
                    value,
                    size_pt=self.theme.body_pt - 1.5,
                    font=self.theme.font_body,
                )

    @staticmethod
    def _column_widths(block: Table, width_cm: float) -> list[float]:
        lengths = []
        for index, column in enumerate(block.columns):
            content_length = max(
                [len(column), *(len(row[index]) for row in block.rows)],
            )
            lengths.append(max(content_length, 6))
        total = sum(lengths) or 1
        return [max(width_cm * length / total, 2.0) for length in lengths]

    def _render_callout(self, document: DocumentObject, block: Callout) -> None:
        if not block.text.strip():
            return
        accent, fill = _CALLOUT_COLORS.get(block.kind, _CALLOUT_COLORS["info"])
        table = document.add_table(rows=1, cols=1)
        table.autofit = False
        cell = table.cell(0, 0)
        cell.width = Cm(16.6)
        set_cell_background(cell, fill)
        set_cell_borders(
            cell,
            left=(accent, 28),
            top=(self.theme.accent_light_hex, 4),
            bottom=(self.theme.accent_light_hex, 4),
            right=(self.theme.accent_light_hex, 4),
        )
        set_cell_margins(cell, top_cm=0.22, bottom_cm=0.18, left_cm=0.35, right_cm=0.25)
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(2)
        format_run(
            paragraph,
            _CALLOUT_LABELS.get(block.kind, "NOTE"),
            bold=True,
            color=accent,
            size_pt=8,
            font=self.theme.font_body,
        )
        paragraph.add_run("\n")
        format_run(
            paragraph,
            block.text,
            color="262626",
            size_pt=self.theme.body_pt - 0.5,
            font=self.theme.font_body,
        )
        spacer = document.add_paragraph()
        spacer.paragraph_format.space_after = Pt(2)
        spacer.paragraph_format.line_spacing = 0.6


def render_spec(spec: DocumentSpec, theme: Theme = DEFAULT_THEME) -> bytes:
    """Convenience function for one-off rendering."""
    return DocxRenderer(theme).render(spec)
