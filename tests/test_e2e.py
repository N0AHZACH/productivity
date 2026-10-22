"""End-to-end HTTP tests with a schema-driven FakeLLM and real .docx rendering."""

import asyncio
import json
import re
from io import BytesIO
from pathlib import Path
from typing import Any

import httpx
import pytest
from docx import Document

from app.agent.composer import CompositionDraft
from app.agent.planner import fallback_plan
from app.agent.reflector import CriticReview
from app.container import Container
from app.core.config import Settings
from app.infra.filestore import FileStore
from app.llm.fake import FakeLLM
from app.main import create_app
from app.models.common import DocType, TaskStatus
from app.models.document import Paragraph, SectionDraft, Table
from app.models.plan import TaskList
from app.models.state import Analysis
from app.render.docx_renderer import DocxRenderer
from app.tools.registry import default_registry


def _standard_analysis() -> Analysis:
    return Analysis(
        doc_type=DocType.PROJECT_PLAN,
        objective="Launch a customer support chatbot for a mid-size retailer.",
        audience="Executive sponsor and delivery team",
        known_facts=["Three-month timeline", "Team of five", "Mid-size retailer"],
        missing_info=[],
        conflicts=[],
        assumptions=["Audience inferred as executive sponsor and delivery team."],
    )


def _complex_analysis() -> Analysis:
    return Analysis(
        doc_type=DocType.DECISION_MEMO,
        objective="Recommend a path for replacing the legacy billing platform.",
        audience="Board of directors",
        known_facts=["Must finish before Q1", "Current contract expires in March"],
        missing_info=["Budget ceiling", "Applicable compliance regime"],
        conflicts=[
            "Cheap, fast, and fully compliant cannot all be maximized together; "
            "prioritize compliance and phase the quick wins."
        ],
        assumptions=[
            "Budget not given; assuming 120,000 USD.",
            "Compliance regime not given; assuming data-protection controls apply.",
        ],
    )


def _section_writer(messages: list[dict[str, str]]) -> SectionDraft:
    user = messages[-1]["content"]
    title_match = re.search(r"^SECTION TITLE: (.+)$", user, flags=re.MULTILINE)
    assert title_match is not None
    title = title_match.group(1).strip()
    analysis_match = re.search(
        r"DOCUMENT ANALYSIS:\n(\{.*?\})\n\nSECTION TITLE:", user, flags=re.DOTALL
    )
    analysis = json.loads(analysis_match.group(1)) if analysis_match else {}
    objective = analysis.get("objective", "the requested objective")
    conflicts = analysis.get("conflicts", [])
    total = re.search(r'"formatted_total":\s*"([^"]+)"', user)
    start = re.search(r'"start_date":\s*"([^"]+)"', user)
    end = re.search(r'"end_date":\s*"([^"]+)"', user)

    opening = (
        f"{title} explains how the work delivers the objective: {objective} "
        "It uses the approved analysis and tool results rather than invented inputs."
    )
    detail = [
        "This section records the decisions, owners, and evidence a reader needs "
        "before approving the next step."
    ]
    if conflicts:
        detail.append("Trade-offs considered: " + " ".join(conflicts))
    if total:
        detail.append(f"The approved budget total is {total.group(1)}.")
    if start and end:
        detail.append(
            f"The delivery schedule runs from {start.group(1)} to {end.group(1)}."
        )
    blocks: list[Any] = [Paragraph(text=opening), Paragraph(text=" ".join(detail))]
    if "budget" in title.lower() and total:
        blocks.append(
            Table(
                columns=["Item", "Amount"],
                rows=[["Approved delivery plan", total.group(1)], ["Total", total.group(1)]],
            )
        )
    return SectionDraft(heading=title, blocks=blocks)


def _fake_llm(analysis: Analysis, composition: CompositionDraft) -> FakeLLM:
    return FakeLLM.by_schema(
        {
            Analysis: analysis,
            TaskList: fallback_plan(analysis),
            SectionDraft: _section_writer,
            CompositionDraft: composition,
            CriticReview: {"passed": True, "issues": []},
        }
    )


def _app(tmp_path: Path, llm: Any, **settings_overrides: Any) -> Any:
    values: dict[str, Any] = {
        "data_dir": tmp_path,
        "reflection_llm_mode": "auto",
        "run_timeout_s": 20,
    }
    values.update(settings_overrides)
    settings = Settings(**values)
    container = Container(
        settings=settings,
        llm=llm,
        registry=default_registry(),
        renderer=DocxRenderer(),
        filestore=FileStore(tmp_path / "documents"),
    )
    return create_app(settings=settings, container=container)


def _docx_text(content: bytes) -> str:
    document = Document(BytesIO(content))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


@pytest.mark.asyncio
async def test_standard_request_end_to_end(tmp_path: Path) -> None:
    analysis = _standard_analysis()
    composition = CompositionDraft(
        title="Customer Support Chatbot Project Plan",
        subtitle="Twelve-week assisted launch",
        executive_summary=(
            "The plan launches a focused assisted support pilot before widening scope. "
            "Success depends on content quality, integration readiness, and measured deflection."
        ),
        recommendation="Start with a narrow pilot and expand only after measured success.",
    )
    llm = _fake_llm(analysis, composition)
    app = _app(tmp_path, llm)
    transport = httpx.ASGITransport(app=app)

    async with httpx.AsyncClient(transport=transport, base_url="http://test") as client:
        response = await client.post(
            "/agent",
            json={
                "request": (
                    "Write a project plan for launching a customer support chatbot at a "
                    "mid-size retailer over three months with a team of five."
                )
            },
        )
        assert response.status_code == 200, response.text
        payload = response.json()
        download = await client.get(payload["document_url"])

    assert payload["status"] == "completed"
    assert payload["plan_source"] == "llm"
    assert payload["document_url"] == f"/documents/{payload['run_id']}"
    assert payload["filename"] == f"{payload['run_id']}.docx"
    assert all(report["passed"] for report in payload["reflection"])
    assert not any(call.schema is CriticReview for call in llm.calls)
    stages = {event["stage"] for event in payload["trace"]}
    assert {"run", "guardrails", "analysis", "plan", "execute", "compose"} <= stages
    assert {"reflection", "render", "store"} <= stages
    assert all(task["status"] == TaskStatus.DONE for task in payload["plan"]["tasks"])

    assert download.status_code == 200
    assert download.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert download.content.startswith(b"PK")
    text = _docx_text(download.content)
    assert "Customer Support Chatbot Project Plan" in text
    assert "Executive Summary" in text
    assert "Budget and Resourcing" in text
    assert "USD" in text
    assert (tmp_path / "documents" / payload["filename"]).read_bytes() == download.content


@pytest.mark.asyncio
async def test_complex_conflict_request_end_to_end(tmp_path: Path) -> None:
    analysis = _complex_analysis()
    composition = CompositionDraft(
        title="Legacy Billing Platform Replacement",
        subtitle="Decision memo for the board",
        executive_summary=(
            "The board must choose a replacement path that balances speed, cost, and "
            "compliance. The recommendation prioritizes compliance and phases the fastest wins."
        ),
        recommendation="Prioritize compliance and phase quick wins after the core controls land.",
    )
    llm = _fake_llm(analysis, composition)
    app = _app(tmp_path, llm)
    transport = httpx.ASGITransport(app=app)

    async with httpx.AsyncClient(transport=transport, base_url="http://test") as client:
        response = await client.post(
            "/agent",
            json={
                "request": (
                    "Prepare a decision memo for the board on replacing our legacy billing "
                    "platform. We need it before Q1, but budget and compliance constraints "
                    "conflict."
                )
            },
        )
        assert response.status_code == 200, response.text
        payload = response.json()
        download = await client.get(payload["document_url"])

    assert payload["status"] == "completed"
    assert payload["analysis"]["doc_type"] == "decision_memo"
    assert payload["conflicts"] == analysis.conflicts
    assert len(payload["assumptions"]) >= 2
    assert payload["reflection"]
    assert payload["reflection"][0]["passed"] is True
    assert any(call.schema is CriticReview for call in llm.calls)
    assert download.status_code == 200
    text = _docx_text(download.content)
    assert "DECISION" in text
    assert "Trade-off:" in text
    assert "Cheap, fast, and fully compliant" in text
    assert "Assumptions and Open Questions" in text


@pytest.mark.asyncio
async def test_document_download_404_and_api_key_auth(tmp_path: Path) -> None:
    llm = FakeLLM()
    app = _app(tmp_path, llm, api_keys="secret-key")
    transport = httpx.ASGITransport(app=app)

    async with httpx.AsyncClient(transport=transport, base_url="http://test") as client:
        unauthorized = await client.post("/agent", json={"request": "Write a plan."})
        missing = await client.get("/documents/" + "a" * 32, headers={"X-API-Key": "secret-key"})
        health = await client.get("/health")

    assert unauthorized.status_code == 401
    assert unauthorized.json()["error"]["code"] == "UNAUTHORIZED"
    assert missing.status_code == 404
    assert missing.json()["error"]["code"] == "NOT_FOUND"
    assert health.status_code == 200


@pytest.mark.asyncio
async def test_out_of_scope_request_returns_standard_error(tmp_path: Path) -> None:
    analysis = Analysis(
        doc_type=DocType.BUSINESS_REPORT,
        objective="Write a poem",
        audience="Unknown",
        in_scope=False,
        refusal_reason="That is a creative-writing request, not a business document.",
    )
    app = _app(tmp_path, FakeLLM.by_schema({Analysis: analysis}))
    transport = httpx.ASGITransport(app=app)

    async with httpx.AsyncClient(transport=transport, base_url="http://test") as client:
        response = await client.post("/agent", json={"request": "Write a poem about clouds."})

    assert response.status_code == 400
    assert response.json()["error"]["code"] == "GUARDRAIL_REJECTED"


class _SlowLLM(FakeLLM):
    async def complete_json(self, *args: Any, **kwargs: Any) -> Any:
        await asyncio.sleep(0.2)
        raise AssertionError("timeout should cancel this call")


class _SlowReflectionLLM(FakeLLM):
    async def complete_json(
        self,
        messages: list[dict[str, str]],
        schema: Any,
        **kwargs: Any,
    ) -> Any:
        if schema is CriticReview:
            await asyncio.sleep(0.2)
        return await super().complete_json(messages, schema, **kwargs)


@pytest.mark.asyncio
async def test_reflection_timeout_renders_current_draft(tmp_path: Path) -> None:
    analysis = _standard_analysis()
    composition = CompositionDraft(
        title="Customer Support Chatbot Project Plan",
        subtitle="Twelve-week assisted launch",
        executive_summary="The plan launches a focused assisted support pilot.",
        recommendation="Start with a narrow pilot and expand after measured success.",
    )
    llm = _SlowReflectionLLM.by_schema(
        {
            Analysis: analysis,
            TaskList: fallback_plan(analysis),
            SectionDraft: _section_writer,
            CompositionDraft: composition,
            CriticReview: {"passed": True, "issues": []},
        }
    )
    app = _app(
        tmp_path,
        llm,
        reflection_llm_mode="always",
        reflection_timeout_s=0.02,
    )
    transport = httpx.ASGITransport(app=app)

    async with httpx.AsyncClient(transport=transport, base_url="http://test") as client:
        response = await client.post(
            "/agent",
            json={"request": "Write a project plan for a support chatbot pilot."},
        )
        assert response.status_code == 200, response.text
        payload = response.json()
        download = await client.get(payload["document_url"])

    assert payload["status"] == "completed_with_warnings"
    assert any("Reflection exceeded its time budget" in item for item in payload["warnings"])
    reflection_events = [
        event for event in payload["trace"] if event["stage"] == "reflection"
    ]
    assert [event["message"] for event in reflection_events] == ["Review timed out"]
    assert download.status_code == 200
    assert download.content.startswith(b"PK")


@pytest.mark.asyncio
async def test_run_timeout_returns_typed_error(tmp_path: Path) -> None:
    app = _app(tmp_path, _SlowLLM(), run_timeout_s=0.02)
    transport = httpx.ASGITransport(app=app)

    async with httpx.AsyncClient(transport=transport, base_url="http://test") as client:
        response = await client.post("/agent", json={"request": "Write a project plan."})

    assert response.status_code == 504
    assert response.json()["error"]["code"] == "RUN_TIMEOUT"
