"""Renderer tests reopen real .docx bytes; no LLM or network."""

import re
from io import BytesIO

import pytest
from docx import Document
from pydantic import ValidationError

from app.models.common import DocType
from app.models.document import (
    BulletList,
    Callout,
    DocumentSpec,
    Heading,
    NumberedList,
    Paragraph,
    Section,
    Table,
)
from app.render.docx_renderer import DocxRenderer
from app.render.themes import Theme


def _spec() -> DocumentSpec:
    return DocumentSpec(
        title="Operations Readiness Review",
        subtitle="Prepared for the executive team",
        doc_type=DocType.BUSINESS_REPORT,
        metadata={
            "author": "DocPilot",
            "date": "2026-10-22",
            "version": "2.0",
        },
        assumptions=["Budget was not provided; assuming 125,000 USD."],
        sections=[
            Section(
                id="s_one",
                heading="Findings",
                blocks=[
                    Paragraph(text="Service demand increased while staffing stayed flat."),
                    Heading(text="Evidence", level=1),
                    BulletList(items=["Contact volume rose 18%", "Headcount was unchanged"]),
                    NumberedList(items=["Measure", "Prioritize", "Act"]),
                ],
            ),
            Section(
                id="s_two",
                heading="Recommendation",
                blocks=[
                    Paragraph(text="Adopt a staged staffing plan."),
                    Table(
                        caption="Investment",
                        columns=["Phase", "Weeks", "Cost"],
                        rows=[["Pilot", "4", "20,000"], ["Scale", "8", "60,000"]],
                    ),
                    NumberedList(items=["Approve pilot", "Review results"]),
                    Callout(kind="decision", text="Proceed with the pilot in November."),
                ],
            ),
        ],
    )


def _render(spec: DocumentSpec) -> tuple[Document, bytes]:
    theme = Theme(footer_text="DocPilot Confidential")
    data = DocxRenderer(theme).render(spec)
    return Document(BytesIO(data)), data


def test_renderer_produces_ordered_headings_and_cover() -> None:
    document, data = _render(_spec())
    assert data.startswith(b"PK")
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    assert "Operations Readiness Review" in paragraphs
    assert "Prepared for the executive team" in paragraphs

    headings = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.style.name in {"Heading 1", "Heading 2", "Heading 3", "Heading 4"}
    ]
    assert headings == [
        "Assumptions and Open Questions",
        "Findings",
        "Evidence",
        "Recommendation",
    ]


def test_table_shape_and_header_furniture() -> None:
    document, _ = _render(_spec())
    table = next(
        table
        for table in document.tables
        if table.rows and [cell.text for cell in table.rows[0].cells] == ["Phase", "Weeks", "Cost"]
    )
    assert len(table.rows) == 3
    assert len(table.columns) == 3

    section = document.sections[0]
    header_text = "\n".join(paragraph.text for paragraph in section.header.paragraphs)
    footer_text = "\n".join(paragraph.text for paragraph in section.footer.paragraphs)
    assert "Operations Readiness Review" in header_text
    assert "DocPilot Confidential" in footer_text
    assert "Page" in footer_text and "of" in footer_text
    assert "PAGE" in section.footer.paragraphs[0]._p.xml
    assert "NUMPAGES" in section.footer.paragraphs[0]._p.xml


def test_document_contains_toc_field_and_repeating_header() -> None:
    document, _ = _render(_spec())
    assert "TOC" in document.element.xml
    content_table = next(
        table
        for table in document.tables
        if table.rows and [cell.text for cell in table.rows[0].cells] == ["Phase", "Weeks", "Cost"]
    )
    assert "tblHeader" in content_table.rows[0]._tr.xml


def test_two_numbered_lists_both_start_at_one() -> None:
    document, _ = _render(_spec())
    starts = [p.text for p in document.paragraphs if re.match(r"^1\.\t", p.text)]
    assert len(starts) == 2


def test_ragged_table_is_rejected_before_render() -> None:
    with pytest.raises(ValidationError, match="column count"):
        Table(columns=["A", "B"], rows=[["only one"]])


def test_empty_strings_are_skipped() -> None:
    spec = DocumentSpec(
        title="",
        doc_type=DocType.SOP,
        sections=[
            Section(
                id="s_empty",
                heading="",
                blocks=[
                    Paragraph(text=""),
                    Heading(text="", level=2),
                    BulletList(items=["", "   "]),
                    NumberedList(items=[""]),
                    Callout(text=""),
                ],
            )
        ],
    )
    document, _ = _render(spec)
    assert "Untitled Document" in [p.text for p in document.paragraphs]


def test_bad_logo_path_is_ignored() -> None:
    theme = Theme(logo_path="definitely/missing/logo.png", footer_text="")
    spec = _spec()
    data = DocxRenderer(theme).render(spec)
    assert data.startswith(b"PK")
